"""
resume_parser.py
Extracts clean text from PDF and DOCX resumes locally using PDFPlumber.
Zero data leaves the machine.
"""

import os
import tempfile
from pathlib import Path

import docx
import pdfplumber
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


SECTION_HEADERS = [
    "experience", "work experience", "employment", "internship",
    "education", "academic", "qualification",
    "skills", "technical skills", "core competencies",
    "projects", "personal projects", "academic projects",
    "certifications", "achievements", "awards",
    "summary", "objective", "profile", "about",
    "activities", "extracurricular", "volunteer",
    "contact", "languages", "publications",
]

MAX_HEADER_CHARS = 50


def _looks_like_section_header(text: str) -> bool:
    """True for short headings such as 'EXPERIENCE' or 'Technical Skills:'."""
    cleaned = text.strip().rstrip(":").strip()
    if not cleaned or len(cleaned) > MAX_HEADER_CHARS:
        return False
    lowered = cleaned.lower()
    if any(h == lowered or lowered.startswith(h + " ") for h in SECTION_HEADERS):
        return True
    letters = [c for c in cleaned if c.isalpha()]
    if not letters:
        return False
    # ALL CAPS headings are the norm in the two-column resume templates.
    return all(c.isupper() for c in letters) and len(letters) > 1


def _is_side_by_side_layout(table) -> bool:
    """True for a 2-column grid whose first row holds two headings.

    Those templates put EXPERIENCE beside TECHNICAL SKILLS, so reading across
    rows would file the left column's bullets under the right column's heading.
    """
    try:
        if len(table.columns) != 2 or not table.rows:
            return False
        first_row = table.rows[0]
        if len(first_row.cells) != 2:
            return False
        left, right = (cell.text.strip() for cell in first_row.cells)
    except (IndexError, ValueError):
        return False
    return _looks_like_section_header(left) and _looks_like_section_header(right)


def _iter_block_items(parent):
    """Yield Paragraph and Table objects in true document order.

    ``Document.paragraphs`` and ``Document.tables`` are two separate lists, so
    walking them in turn scrambles the reading order of grid layouts: a
    two-column resume came out as "all prose, then all table cells", which made
    section detection attribute experience text to the wrong section.
    """
    if isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        parent_element = parent.element.body

    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF resume using PDFPlumber (100% local)."""
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX resume, including tables, in document order."""
    document = docx.Document(file_path)
    parts = []

    def emit(text: str) -> None:
        text = text.strip()
        if text:
            parts.append(text)

    def walk_table(table) -> None:
        if _is_side_by_side_layout(table):
            # Two columns that both start with a heading are a side-by-side
            # layout: read the left column top-to-bottom, then the right one,
            # so each column's content stays under its own heading.
            for column in range(len(table.columns)):
                emit_column(table, column)
        else:
            for row in table.rows:
                emit_row(row)

    def emit_column(table, column: int) -> None:
        for row in table.rows:
            cells = row.cells
            if column < len(cells):
                emit_cell(cells[column])

    def emit_row(row) -> None:
        emitted = []
        for cell in row.cells:
            # A merged cell is reported once per grid position it spans.
            element = cell._tc
            if any(element is seen for seen in emitted):
                continue
            emitted.append(element)
            emit_cell(cell)

    def emit_cell(cell) -> None:
        for block in _iter_block_items(cell):
            if isinstance(block, Paragraph):
                emit(block.text)
            else:
                walk_table(block)

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            emit(block.text)
        else:
            walk_table(block)

    return "\n".join(parts)


def parse_resume(uploaded_file) -> dict:
    """
    Parse uploaded resume file (PDF or DOCX).
    Returns a dict with raw_text and structured sections.
    """
    file_name = uploaded_file.name
    extension = Path(file_name).suffix.lower()

    # Save temp file
    with tempfile.NamedTemporaryFile(delete=False, suffix=extension) as tmp:
        tmp.write(uploaded_file.read())
        tmp_path = tmp.name

    try:
        if extension == ".pdf":
            raw_text = extract_text_from_pdf(tmp_path)
        elif extension in [".docx", ".doc"]:
            raw_text = extract_text_from_docx(tmp_path)
        else:
            raise ValueError(f"Unsupported file type: {extension}")
    finally:
        os.unlink(tmp_path)

    sections = extract_sections(raw_text)
    word_count = len(raw_text.split())
    line_count = len([l for l in raw_text.split("\n") if l.strip()])

    return {
        "raw_text": raw_text,
        "sections": sections,
        "word_count": word_count,
        "line_count": line_count,
        "file_name": file_name,
    }


def extract_sections(text: str) -> dict:
    """
    Heuristically detect resume sections (Experience, Education, Skills, Projects, etc.)
    """
    section_headers = SECTION_HEADERS

    lines = text.split("\n")
    sections = {}
    current_section = "header"
    current_lines = []

    for line in lines:
        stripped = line.strip()
        lower = stripped.lower()

        is_header = any(
            lower == h or lower.startswith(h + " ") or lower.startswith(h + ":")
            for h in section_headers
        )

        if is_header and len(stripped) < MAX_HEADER_CHARS:
            if current_lines:
                sections[current_section] = "\n".join(current_lines).strip()
            # "Languages: Python, SQL" is both a heading and a content line, so
            # keep the tail instead of discarding it.
            heading, _, tail = stripped.partition(":")
            current_section = heading.lower().strip()
            current_lines = [tail.strip()] if tail.strip() else []
        else:
            if stripped:
                current_lines.append(stripped)

    if current_lines:
        sections[current_section] = "\n".join(current_lines).strip()

    return sections
